"""2022-2024 gap test (docs/gap2022_preregistration.md): the live model graded against bet365 opening lines from
the AusSportsBetting nfl.xlsx (kept out of git). Also the pre-grading checks, including a 2018-2021 proxy check."""
from __future__ import annotations

import datetime as dt
import math

import numpy as np
import polars as pl

from engine import config
from engine.backtest import total as T
from engine.picks.grading import is_half_point
from engine.teams import to_canonical

XLSX = config.ROOT / "data" / "raw" / "aussportsbetting" / "nfl.xlsx"
FR = {"OAK": "LV", "SD": "LAC", "STL": "LA"}


def load_xlsx(path=XLSX) -> pl.DataFrame:
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Data"]
    rows = ws.iter_rows(values_only=True)
    head = next(rows)
    ix = {h: i for i, h in enumerate(head)}
    out = []
    for r in rows:
        if not r[ix["Date"]]:
            continue
        h, a = (to_canonical(r[ix["Home Team"]]), to_canonical(r[ix["Away Team"]]))
        day = r[ix["Date"]]
        day = dt.date.fromisoformat(day[:10]) if isinstance(day, str) else day.date()
        out.append({"date": day, "home": FR.get(h, h), "away": FR.get(a, a),
                    "open": r[ix["Home Line Open"]], "close_b": r[ix["Home Line Close"]],
                    "min_b": r[ix["Home Line Min"]], "max_b": r[ix["Home Line Max"]],
                    "hp_close": r[ix["Home Line Odds Close"]], "ap_close": r[ix["Away Line Odds Close"]],
                    "hml_close": r[ix["Home Odds Close"]], "aml_close": r[ix["Away Odds Close"]],
                    "neutral": r[ix["Neutral Venue?"]] == "Y", "playoff": r[ix["Playoff Game?"]] == "Y"})
    return pl.DataFrame(out, infer_schema_length=None)


def match(x: pl.DataFrame, sched: pl.DataFrame) -> pl.DataFrame:
    """nflverse game_id -> bet365 open/close as the nflverse HOME spread (flipped when the file lists the teams the
    other way round, e.g. neutral sites). Date within +-1 day."""
    by = {}
    for r in x.iter_rows(named=True):
        by.setdefault(frozenset((r["home"], r["away"])), []).append(r)
    out = []
    for g in sched.iter_rows(named=True):
        day = dt.date.fromisoformat(g["gameday"])
        h, a = FR.get(g["home_team"], g["home_team"]), FR.get(g["away_team"], g["away_team"])
        c = [r for r in by.get(frozenset((h, a)), []) if abs((r["date"] - day).days) <= 1]
        if len(c) != 1:
            continue
        r = c[0]
        s = 1.0 if r["home"] == h else -1.0
        f = lambda v: None if v is None else s * float(v)
        sw = s < 0
        lo, hi = f(r["min_b"]), f(r["max_b"])
        out.append({"game_id": g["game_id"], "b_open": f(r["open"]), "b_close": f(r["close_b"]),
                    "b_min": None if lo is None else min(lo, hi), "b_max": None if hi is None else max(lo, hi),
                    "hp": r["ap_close"] if sw else r["hp_close"], "ap": r["hp_close"] if sw else r["ap_close"],
                    "hml": r["aml_close"] if sw else r["hml_close"], "aml": r["hml_close"] if sw else r["aml_close"],
                    "swapped": sw})
    return pl.DataFrame(out, infer_schema_length=None)


def _rec(w, n):
    r = w / n if n else float("nan")
    return {"W": int(w), "L": int(n - w), "N": int(n), "win_rate": r, "se": math.sqrt(r * (1 - r) / n) if n else None}


def _grade(tr, te, col):
    te = te.filter(pl.col(col).is_not_null(), pl.col("margin") != -pl.col(col))
    p, _ = T._predict("shifted_total", tr, te, col)
    y = te["margin"].to_numpy() > -te[col].to_numpy()
    return te.with_columns(pl.Series("p_home", p), pl.Series("win", (p > 0.5) == y))


def run() -> dict:
    df, sched = T.load()               # nflverse REG with close, total; SBRO open (2007-2021)
    reg = sched.filter(pl.col("game_type") == "REG")
    m = match(load_xlsx(), reg.filter(pl.col("season").is_between(2018, 2024)))
    d = df.join(m, on="game_id", how="left").join(reg.select("game_id", "week"), on="game_id")
    out = {"checks": {}, "proxy_2018_2021": {}, "seasons": {}}
    for s in (2022, 2023, 2024):
        g = d.filter(pl.col("season") == s)
        o = g.filter(pl.col("b_open").is_not_null())
        both = o.filter(pl.col("b_close").is_not_null())
        out["checks"][s] = {"games": g.height, "with_open": o.height,
                            "whole_number_opens": int(sum(not is_half_point(v) for v in o["b_open"])),
                            "close_within_half_pt": float(((both["b_close"] - both["close"]).abs() <= 0.5).mean()),
                            "open_eq_close": float((o["b_open"] == o["close"]).mean())}
    # proxy check: same model, same seasons, bet365 open vs SBRO open
    for s in range(2018, 2022):
        g = d.filter(pl.col("season") == s)
        both = g.filter(pl.col("b_open").is_not_null(), pl.col("open").is_not_null())
        tr = df.filter(pl.col("season") < s)
        gb, gs = _grade(tr, g, "b_open"), _grade(tr, g, "open")
        out["proxy_2018_2021"][s] = {"open_equal": float((both["b_open"] == both["open"]).mean()),
                                     "open_within_1": float(((both["b_open"] - both["open"]).abs() <= 1).mean()),
                                     "bet365": _rec(gb["win"].sum(), gb.height), "sbro": _rec(gs["win"].sum(), gs.height)}
    games = []
    for s in (2022, 2023, 2024):
        g = d.filter(pl.col("season") == s)
        games.append(_grade(df.filter(pl.col("season") < s), g, "b_open"))
    G = pl.concat(games).with_columns((pl.col("close") - pl.col("b_open")).alias("drift"),
                                      pl.col("b_open").map_elements(is_half_point, return_dtype=pl.Boolean).alias("half"))
    G.select("game_id", "season", "week", "b_open", "close", "margin", "p_home", "win", "drift", "half") \
        .write_csv(config.ROOT / "data" / "backtests" / "gap2022_games.csv")
    H = G.filter(pl.col("half"))
    out["primary_half_point"] = _rec(H["win"].sum(), H.height)
    for s in (2022, 2023, 2024):
        x, xh = G.filter(pl.col("season") == s), H.filter(pl.col("season") == s)
        out["seasons"][s] = {"half_point": _rec(xh["win"].sum(), xh.height), "all": _rec(x["win"].sum(), x.height)}
    out["all_games"] = _rec(G["win"].sum(), G.height)
    mv = G.filter(pl.col("drift").abs() >= 1)
    out["moved_1plus"] = _rec(mv["win"].sum(), mv.height)
    nm = G.filter(pl.col("drift") == 0)
    out["no_move"] = _rec(nm["win"].sum(), nm.height)
    wk = G.group_by("season", "week").agg(pl.col("win").sum().alias("W"))
    out["weeks_12plus"] = {s: int(wk.filter(pl.col("season") == s, pl.col("W") >= 12).height) for s in (2022, 2023, 2024)}
    r, se = out["primary_half_point"]["win_rate"], out["primary_half_point"]["se"]
    lo, hi = r - 1.96 * se, r + 1.96 * se
    out["verdict"] = ("consistent with the 2007-2021 edge" if lo <= 0.549 <= hi and lo > 0.5 else
                      "consistent with no edge" if lo <= 0.5 <= hi and hi < 0.549 else "inconclusive")
    out["ci95"] = (lo, hi)
    return out
